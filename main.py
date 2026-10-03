import asyncio
import json
import re
from urllib.parse import urljoin

from playwright.async_api import async_playwright
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


# ============================================================
# 1. CẤU HÌNH
# ============================================================

BASE_URL = "https://fptshop.com.vn"
PHONE_URL = f"{BASE_URL}/dien-thoai"

OUTPUT_FILE = "fptshop_dien_thoai.xlsx"


# ============================================================
# 2. CHUẨN HÓA TEXT
# ============================================================

def clean_text(text):
    if not text:
        return ""

    text = text.replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text)

    return text.strip()


# ============================================================
# 3. CLICK XEM THÊM CHO ĐẾN KHI HẾT
# ============================================================

async def load_all_products(page):

    print()
    print("=" * 70)
    print("ĐANG TẢI TOÀN BỘ ĐIỆN THOẠI")
    print("=" * 70)

    click_count = 0

    while True:

        try:

            buttons = page.get_by_role(
                "button",
                name=re.compile(
                    r"Xem thêm",
                    re.I
                )
            )

            button_count = await buttons.count()

            visible_button = None

            # Tìm nút Xem thêm đang hiển thị
            for i in range(button_count):

                try:

                    button = buttons.nth(i)

                    if not await button.is_visible():
                        continue

                    if await button.is_disabled():
                        continue

                    visible_button = button
                    break

                except:
                    continue

            # Không còn nút
            if visible_button is None:

                print()
                print("✓ KHÔNG CÒN NÚT XEM THÊM")
                break

            # Cuộn tới nút
            await visible_button.scroll_into_view_if_needed()

            await page.wait_for_timeout(1000)

            # Click
            await visible_button.click(
                timeout=15000
            )

            click_count += 1

            print(
                f"→ Click Xem thêm lần {click_count}"
            )

            # Chờ FPT load sản phẩm
            await page.wait_for_timeout(4000)

            # Scroll nhẹ
            await page.mouse.wheel(
                0,
                800
            )

            await page.wait_for_timeout(1000)

        except Exception as e:

            print(
                f"⚠ Lỗi click: {e}"
            )

            # Không dừng ngay
            await page.wait_for_timeout(3000)

            # Kiểm tra lại xem còn button không
            try:

                buttons = page.get_by_role(
                    "button",
                    name=re.compile(
                        r"Xem thêm",
                        re.I
                    )
                )

                count = await buttons.count()

                if count == 0:

                    print(
                        "✓ Không còn nút Xem thêm."
                    )

                    break

            except:

                pass

    print()
    print(
        f"ĐÃ CLICK XEM THÊM: {click_count} LẦN"
    )


# ============================================================
# 4. LẤY LINK SẢN PHẨM
# ============================================================

async def get_product_links(page):

    print()
    print("=" * 70)
    print("ĐANG LẤY LINK SẢN PHẨM")
    print("=" * 70)

    # Scroll toàn bộ trang
    await page.evaluate("""
        async () => {

            await new Promise(resolve => {

                let current = 0;

                const timer = setInterval(() => {

                    window.scrollBy(
                        0,
                        700
                    );

                    current += 700;

                    if (
                        current >=
                        document.body.scrollHeight
                    ) {

                        clearInterval(timer);
                        resolve();

                    }

                }, 150);

            });

        }
    """)

    await page.wait_for_timeout(3000)

    # Lấy toàn bộ a
    raw_links = await page.locator(
        "a[href]"
    ).evaluate_all("""
        els => els.map(a => ({
            href: a.href,
            text: a.innerText.trim()
        }))
    """)

    result = []

    # Các URL hãng/category cần loại
    category_slugs = {
        "apple",
        "samsung",
        "xiaomi",
        "oppo",
        "vivo",
        "realme",
        "tecno",
        "nokia",
        "honor",
        "masstel",
        "google",
        "sony",
        "oneplus",
        "asus",
        "zte",
        "motorola"
    }

    for item in raw_links:

        href = item.get(
            "href",
            ""
        )

        if not href:
            continue

        href = urljoin(
            BASE_URL,
            href
        )

        href = href.split("?")[0]
        href = href.rstrip("/")

        # Chỉ URL điện thoại
        if "/dien-thoai/" not in href:
            continue

        # Bỏ trang danh mục
        if href == PHONE_URL:
            continue

        slug = href.split(
            "/dien-thoai/",
            1
        )[-1]

        if not slug:
            continue

        # Bỏ category
        if slug.lower() in category_slugs:
            continue

        # Chống trùng
        if href not in result:

            result.append(href)

    print()
    print(
        f"→ TÌM ĐƯỢC {len(result)} LINK SẢN PHẨM"
    )

    return result


# ============================================================
# 5. LẤY JSON-LD
# ============================================================

async def get_json_ld(page):

    data_list = []

    scripts = page.locator(
        'script[type="application/ld+json"]'
    )

    count = await scripts.count()

    for i in range(count):

        try:

            raw = await scripts.nth(
                i
            ).text_content()

            if not raw:
                continue

            data = json.loads(
                raw.strip()
            )

            if isinstance(
                data,
                list
            ):

                data_list.extend(data)

            elif isinstance(
                data,
                dict
            ):

                data_list.append(data)

        except:

            continue

    return data_list


# ============================================================
# 6. PRODUCT JSON-LD
# ============================================================

def get_product_jsonld(data_list):

    for item in data_list:

        if not isinstance(
            item,
            dict
        ):
            continue

        if item.get(
            "@type"
        ) == "Product":

            return item

        graph = item.get(
            "@graph"
        )

        if isinstance(
            graph,
            list
        ):

            for obj in graph:

                if (
                    isinstance(
                        obj,
                        dict
                    )
                    and
                    obj.get(
                        "@type"
                    ) == "Product"
                ):

                    return obj

    return {}


# ============================================================
# 7. TÊN SẢN PHẨM
# ============================================================

async def get_product_name(
    page,
    product
):

    # Ưu tiên JSON-LD
    if product:

        name = clean_text(
            product.get(
                "name",
                ""
            )
        )

        if name:
            return name

    # Fallback h1
    try:

        h1 = page.locator(
            "h1"
        ).first

        if await h1.count():

            name = await h1.inner_text()

            return clean_text(
                name
            )

    except:

        pass

    return ""


# ============================================================
# 8. FORMAT GIÁ
# ============================================================

def format_price(value):

    if value is None:
        return ""

    try:

        number = float(
            str(value)
            .replace(",", "")
        )

        return (
            f"{int(number):,}"
            .replace(",", ".")
            + "đ"
        )

    except:

        return ""


# ============================================================
# 9. LẤY GIÁ
# ============================================================

async def get_prices(
    page,
    product
):

    sale_price = ""
    old_price = ""

    # --------------------------------------------------------
    # Giá bán từ JSON-LD
    # --------------------------------------------------------

    if product:

        offers = product.get(
            "offers"
        )

        if isinstance(
            offers,
            list
        ):

            offers = (
                offers[0]
                if offers
                else {}
            )

        if isinstance(
            offers,
            dict
        ):

            sale_price = format_price(
                offers.get(
                    "price"
                )
            )

    # --------------------------------------------------------
    # Lấy text trang
    # --------------------------------------------------------

    try:

        body = await page.locator(
            "body"
        ).inner_text()

        body = clean_text(
            body
        )

    except:

        body = ""

    # Chỉ xét vùng đầu trang
    first_part = body[:10000]

    prices = re.findall(
        r"\d{1,3}(?:[.,]\d{3})+\s*đ",
        first_part
    )

    prices = [
        x.replace(
            " ",
            ""
        )
        for x in prices
    ]

    unique_prices = []

    for price in prices:

        if price not in unique_prices:

            unique_prices.append(
                price
            )

    # --------------------------------------------------------
    # Fallback giá bán
    # --------------------------------------------------------

    if not sale_price:

        for price in unique_prices:

            sale_price = price
            break

    # --------------------------------------------------------
    # Giá gốc
    # --------------------------------------------------------

    if sale_price:

        try:

            sale_number = int(
                sale_price
                .replace(
                    ".",
                    ""
                )
                .replace(
                    "đ",
                    ""
                )
            )

            candidates = []

            for price in unique_prices:

                try:

                    number = int(
                        price
                        .replace(
                            ".",
                            ""
                        )
                        .replace(
                            "đ",
                            ""
                        )
                    )

                    if number > sale_number:

                        candidates.append(
                            (
                                number,
                                price
                            )
                        )

                except:

                    continue

            if candidates:

                candidates.sort(
                    key=lambda x: x[0]
                )

                old_price = candidates[0][1]

        except:

            pass

    return (
        old_price,
        sale_price
    )


# ============================================================
# 10. BODY LINES
# ============================================================

async def get_body_lines(page):

    try:

        text = await page.locator(
            "body"
        ).inner_text()

        lines = []

        for line in text.splitlines():

            line = clean_text(
                line
            )

            if line:
                lines.append(
                    line
                )

        return lines

    except:

        return []


# ============================================================
# 11. THÔNG SỐ
# ============================================================

def extract_specs(lines):

    specs = {}

    labels = [
        "Chip xử lý (CPU)",
        "Kích thước màn hình",
        "Hệ điều hành",
        "Dung lượng RAM",
        "Bộ nhớ trong",
        "Dung lượng pin",
        "Camera sau",
        "Camera trước",
        "Tần số quét"
    ]

    for i, line in enumerate(
        lines
    ):

        current = clean_text(
            line
        )

        for label in labels:

            # label: value
            if current.lower().startswith(
                label.lower() + ":"
            ):

                value = current[
                    current.find(":") + 1:
                ]

                value = clean_text(
                    value
                )

                if value:

                    specs[label] = value

            # label nằm riêng một dòng
            elif (
                current.lower()
                == label.lower()
            ):

                if (
                    i + 1
                    <
                    len(lines)
                ):

                    value = clean_text(
                        lines[i + 1]
                    )

                    if value:

                        specs[label] = value

    return specs


# ============================================================
# 12. MÔ TẢ
# ============================================================

def extract_description(
    lines
):

    start = -1

    for i, line in enumerate(
        lines
    ):

        if line.lower() == (
            "mô tả sản phẩm"
        ):

            start = i
            break

    if start == -1:
        return ""

    result = []

    stop_words = [
        "đánh giá và bình luận",
        "sản phẩm liên quan",
        "câu hỏi thường gặp"
    ]

    for line in lines[
        start + 1:
    ]:

        lower = line.lower()

        if any(
            word in lower
            for word in stop_words
        ):

            break

        if lower in [
            "xem thêm",
            "thu gọn"
        ]:

            continue

        result.append(
            line
        )

        if len(
            " ".join(result)
        ) > 5000:

            break

    return clean_text(
        " ".join(result)
    )


# ============================================================
# 13. FORMAT EXCEL
# ============================================================

def format_excel(ws):

    # Header
    for cell in ws[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    ws.freeze_panes = "A2"

    ws.auto_filter.ref = (
        ws.dimensions
    )

    widths = {
        "A": 40,
        "B": 18,
        "C": 18,
        "D": 20,
        "E": 30,
        "F": 25,
        "G": 100,
        "H": 100,
        "I": 80
    }

    for col, width in widths.items():

        ws.column_dimensions[
            col
        ].width = width

    for row in ws.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


# ============================================================
# 14. MAIN
# ============================================================

async def main():

    async with async_playwright() as p:

        browser = await p.chromium.launch(
            headless=False
        )

        page = await browser.new_page(
            viewport={
                "width": 1440,
                "height": 900
            }
        )

        # ====================================================
        # MỞ FPT
        # ====================================================

        print()
        print(
            "[1] MỞ FPT SHOP..."
        )

        await page.goto(
            PHONE_URL,
            wait_until="domcontentloaded",
            timeout=60000
        )

        await page.wait_for_timeout(
            5000
        )

        # ====================================================
        # CLICK XEM THÊM
        # ====================================================

        await load_all_products(
            page
        )

        # ====================================================
        # LẤY LINK SAU KHI LOAD XONG
        # ====================================================

        product_links = (
            await get_product_links(
                page
            )
        )

        print()
        print(
            "=" * 70
        )
        print(
            f"TỔNG SẢN PHẨM TÌM ĐƯỢC: {len(product_links)}"
        )
        print(
            "=" * 70
        )

        if len(product_links) == 0:

            print(
                "❌ KHÔNG TÌM ĐƯỢC SẢN PHẨM"
            )

            await browser.close()
            return

        # ====================================================
        # TẠO EXCEL
        # ====================================================

        wb = Workbook()

        ws = wb.active

        ws.title = "DienThoai"

        ws.append([
            "Tên sản phẩm",
            "Giá gốc",
            "Giá bán",
            "Màn hình",
            "Chip",
            "Hệ điều hành",
            "Thông số",
            "Mô tả",
            "URL"
        ])

        # ====================================================
        # CÀO CHI TIẾT
        # ====================================================

        print()
        print(
            "=" * 70
        )
        print(
            "BẮT ĐẦU CÀO CHI TIẾT"
        )
        print(
            "=" * 70
        )

        for index, url in enumerate(
            product_links,
            1
        ):

            print()
            print(
                "-" * 70
            )

            print(
                f"[{index}/{len(product_links)}]"
            )

            print(
                url
            )

            try:

                # --------------------------------------------
                # Mở trang sản phẩm
                # --------------------------------------------

                await page.goto(
                    url,
                    wait_until="domcontentloaded",
                    timeout=60000
                )

                await page.wait_for_timeout(
                    2000
                )

                # --------------------------------------------
                # JSON-LD
                # --------------------------------------------

                jsonld = await get_json_ld(
                    page
                )

                product = (
                    get_product_jsonld(
                        jsonld
                    )
                )

                # --------------------------------------------
                # Tên
                # --------------------------------------------

                name = (
                    await get_product_name(
                        page,
                        product
                    )
                )

                # --------------------------------------------
                # Giá
                # --------------------------------------------

                old_price, sale_price = (
                    await get_prices(
                        page,
                        product
                    )
                )

                # --------------------------------------------
                # Lines
                # --------------------------------------------

                lines = (
                    await get_body_lines(
                        page
                    )
                )

                # --------------------------------------------
                # Specs
                # --------------------------------------------

                specs = extract_specs(
                    lines
                )

                screen = specs.get(
                    "Kích thước màn hình",
                    ""
                )

                chip = specs.get(
                    "Chip xử lý (CPU)",
                    ""
                )

                operating_system = (
                    specs.get(
                        "Hệ điều hành",
                        ""
                    )
                )

                specs_text = " | ".join(
                    f"{key}: {value}"
                    for key, value
                    in specs.items()
                )

                # --------------------------------------------
                # Description
                # --------------------------------------------

                description = (
                    extract_description(
                        lines
                    )
                )

                if not description:

                    description = clean_text(
                        product.get(
                            "description",
                            ""
                        )
                    )

                # --------------------------------------------
                # In terminal
                # --------------------------------------------

                print(
                    f"Tên: {name}"
                )

                print(
                    f"Giá gốc: {old_price}"
                )

                print(
                    f"Giá bán: {sale_price}"
                )

                print(
                    f"Màn hình: {screen}"
                )

                print(
                    f"Chip: {chip}"
                )

                print(
                    f"Hệ điều hành: {operating_system}"
                )

                # --------------------------------------------
                # Ghi Excel
                # --------------------------------------------

                ws.append([
                    name,
                    old_price,
                    sale_price,
                    screen,
                    chip,
                    operating_system,
                    specs_text,
                    description,
                    url
                ])

                format_excel(
                    ws
                )

                # Lưu ngay
                wb.save(
                    OUTPUT_FILE
                )

                print(
                    "✓ Đã lưu Excel"
                )

            except Exception as e:

                print(
                    f"✗ LỖI: {e}"
                )

                # Ghi URL lỗi để không mất sản phẩm
                ws.append([
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    f"LỖI: {e}",
                    url
                ])

                format_excel(
                    ws
                )

                wb.save(
                    OUTPUT_FILE
                )

        # ====================================================
        # HOÀN TẤT
        # ====================================================

        format_excel(
            ws
        )

        wb.save(
            OUTPUT_FILE
        )

        print()
        print(
            "=" * 70
        )

        print(
            "🎉 CÀO HOÀN TẤT"
        )

        print(
            f"Tổng sản phẩm: {len(product_links)}"
        )

        print(
            f"File Excel: {OUTPUT_FILE}"
        )

        print(
            "=" * 70
        )

        await browser.close()


# ============================================================
# CHẠY
# ============================================================

if __name__ == "__main__":

    asyncio.run(main())